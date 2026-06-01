"""
Парсер DOCX файлов (Microsoft Word).
Извлекает текст и таблицы из документов.
Соответствует ISO 27001 A.12.4 (Логирование).
"""

import logging
from typing import Dict, List, Any, Optional
from pathlib import Path
from docx import Document
from docx.table import Table

logger = logging.getLogger(__name__)


class DocxParser:
    """Парсер для DOCX файлов."""

    def __init__(self):
        self.supported_extensions = ['.docx']

    def parse(self, file_path: str) -> Dict[str, Any]:
        """
        Парсит DOCX файл.
        Извлекает таблицы и текст.
        
        Args:
            file_path: Путь к файлу
            
        Returns:
            Dict с headers, rows и errors
        """
        try:
            logger.info(f"Чтение DOCX файла: {file_path}")
            
            doc = Document(file_path)
            
            # Пробуем найти таблицы в документе
            tables_data = self._extract_tables(doc)
            
            if tables_data:
                # Если есть таблицы, используем первую непустую
                for table_data in tables_data:
                    if table_data['rows']:
                        logger.info(f"Найдена таблица: {table_data['row_count']} строк")
                        return {
                            'headers': table_data['headers'],
                            'rows': table_data['rows'],
                            'source_type': 'table',
                            'total_rows': table_data['row_count'],
                            'errors': []
                        }
            
            # Если таблиц нет или они пустые, извлекаем текст
            text_data = self._extract_text(doc)
            
            if text_data['lines']:
                logger.info(f"Извлечен текст: {len(text_data['lines'])} строк")
                return {
                    'headers': ['Текст'],
                    'rows': [[line] for line in text_data['lines']],
                    'source_type': 'text',
                    'total_rows': len(text_data['lines']),
                    'errors': text_data.get('warnings', [])
                }
            
            # Документ пуст
            return {
                'headers': [],
                'rows': [],
                'source_type': 'empty',
                'total_rows': 0,
                'errors': ['Документ не содержит таблиц или текста']
            }
            
        except Exception as e:
            error_msg = f"Ошибка при чтении DOCX: {str(e)}"
            logger.error(error_msg)
            return {
                'headers': [],
                'rows': [],
                'errors': [error_msg]
            }

    def _extract_tables(self, doc: Document) -> List[Dict[str, Any]]:
        """Извлекает все таблицы из документа."""
        tables_data = []
        
        for table in doc.tables:
            table_data = self._parse_table(table)
            if table_data:
                tables_data.append(table_data)
        
        return tables_data

    def _parse_table(self, table: Table) -> Optional[Dict[str, Any]]:
        """Парсит одну таблицу."""
        try:
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    text = cell.text.strip()
                    cells.append(text)
                
                # Пропускаем полностью пустые строки
                if any(cell != '' for cell in cells):
                    rows.append(cells)
            
            if not rows:
                return None
            
            # Первая строка - заголовки
            headers = rows[0] if rows else []
            data_rows = rows[1:] if len(rows) > 1 else []
            
            return {
                'headers': headers,
                'rows': data_rows,
                'row_count': len(data_rows)
            }
            
        except Exception as e:
            logger.warning(f"Ошибка при парсинге таблицы: {str(e)}")
            return None

    def _extract_text(self, doc: Document) -> Dict[str, Any]:
        """Извлекает текст из документа построчно."""
        lines = []
        warnings = []
        
        # Извлекаем параграфы
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # Разбиваем на строки по переносам
                sub_lines = text.split('\n')
                lines.extend([line.strip() for line in sub_lines if line.strip()])
        
        # Предупреждение если документ большой
        if len(lines) > 1000:
            warnings.append(f"Большой документ: {len(lines)} строк. Возможны проблемы с производительностью.")
        
        return {
            'lines': lines,
            'warnings': warnings
        }
